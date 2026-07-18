# [兼容修复] 必须在任何间接依赖 torchvision 的库（FlagEmbedding / transformers 等）之前，
# 先干净地导入 torch / torchvision，使其内部 meta 算子（如 roi_align）一次性注册并缓存，
# 避免后续被重复注册而抛出 "already a kernel registered ... roi_align" 的 RuntimeError
# （torch 2.11 + torchvision 0.26 在 transformers 间接导入时偶发）。
import torch  # noqa: F401
import torchvision  # noqa: F401

import os
import shutil
import sys
import uuid
import json
import time
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any

# [修复] 让「直接运行本文件」也能正确导入项目内的 config/processor/utils 等包。
# 直接执行 python web/api/import_service.py 时，Python 只会把脚本所在目录 web/api 加入 sys.path，
# 而 config、processor、utils、tool 都在项目根目录下，于是会抛
# ModuleNotFoundError: No module named 'config.minio_config'（服务根本起不来、端口不监听）。
# 这里在导入项目内模块之前，把项目根目录显式加入 sys.path，使直接运行与 python main.py 行为一致。
_PROJECT_ROOT = Path(__file__).resolve().parents[2]  # web/api/import_service.py -> 项目根目录
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from starlette.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from config.minio_config import minio_config
from processor.import_processor.main_graph import KBImportWorkflow
from tool.logger import logger
import uvicorn
from fastapi import FastAPI, HTTPException, BackgroundTasks, UploadFile, File
from starlette.middleware.cors import CORSMiddleware

from utils.minio_utils import get_minio_client
from utils.task_utils import update_task_status, add_done_task, add_running_task, get_task_status, get_done_task_list, \
    get_running_task_list, set_task_result, get_task_result

# 1. 创建应用
# 标题和描述会在Swagger文档中展示
app = FastAPI(
    title="鹏至智库-导入API",
    description="此文档是鹏至智库导入流程的API接口说明"
)

# [安全加固 Phase0] 跨域收紧：禁止 `*` + credentials 的非法组合，
# 改为显式来源白名单（环境变量 KB_CORS_ORIGINS 可覆盖，逗号分隔）。
# 原 allow_origins=["*"] + allow_credentials=True 既会被浏览器拒绝，
# 又意味着任何网站都能代表用户发起带凭证请求，属高危误配置。
def _kb_cors_origins() -> list:
    raw = os.getenv("KB_CORS_ORIGINS", "")
    if raw:
        return [o.strip() for o in raw.split(",") if o.strip()]
    # 默认仅放行同源的导入/问答前端（前端由本服务同源托管，一般无需跨域）
    return [
        "http://127.0.0.1:8100", "http://127.0.0.1:8101",
        "http://localhost:8100", "http://localhost:8101",
    ]


# 2. 跨域
app.add_middleware(
    CORSMiddleware,
    allow_origins=_kb_cors_origins(),  # 显式白名单
    allow_credentials=False,  # 同源无需凭证，关闭以根除凭证泄露面
    allow_methods=["*"],  # 允许的请求方法
    allow_headers=["*"],  # 允许的请求头
)


# [安全加固 Phase0] 可选 API 密钥闸门（防御纵深）。
# 设置环境变量 KB_API_KEY 后，/upload、/stats 等变更/敏感接口必须携带
# `X-API-Key: <key>` 或 `Authorization: Bearer <key>`，否则返回 401。
# 未设置时保持开发态开放（服务已默认绑定 127.0.0.1），但启动会打印醒目告警，
# 提醒生产环境务必配置强随机 KB_API_KEY。
_KB_API_KEY = os.getenv("KB_API_KEY", "").strip()
if _KB_API_KEY:
    from starlette.middleware.base import BaseHTTPMiddleware
    from starlette.responses import JSONResponse

    class _ApiKeyGuard(BaseHTTPMiddleware):
        async def dispatch(self, request, call_next):
            if request.url.path in ("/upload", "/stats"):
                auth = request.headers.get("Authorization", "")
                token = auth[7:].strip() if auth.lower().startswith("bearer ") else request.headers.get("X-API-Key", "").strip()
                if token != _KB_API_KEY:
                    return JSONResponse(
                        status_code=401,
                        content={"ok": False, "detail": "缺少或错误的 API 密钥"},
                    )
            return await call_next(request)

    app.add_middleware(_ApiKeyGuard)
else:
    import logging as _logging
    _logging.getLogger("knowledge_base").warning(
        "KB_API_KEY 未设置：导入/统计接口当前无任何鉴权（仅依赖 127.0.0.1 绑定）。"
        "生产环境请设置强随机 KB_API_KEY 并在调用方携带 X-API-Key。"
    )


# [安全加固 Phase0] 上传文件名净化：防止 `../../etc/passwd` 之类路径穿越。
# 仅保留基础文件名，剥离任何目录成分与不可见字符。
def _safe_filename(name: str) -> str:
    name = os.path.basename(name or "").strip()
    name = name.replace("\\", "_").replace("/", "_").replace("\x00", "")
    return name or "unnamed.bin"


# 2.5 健康检查：用于服务探活（与 query_service 对齐，方便一键启动后确认导入服务已就绪）
@app.get("/health")
async def health():
    """检查导入服务是否正常"""
    return {"ok": True}


# 2.6 知识库统计：聚合 doc/ 目录产物 + 内存态任务状态，返回真实运行指标（供概览页使用）
_STATS_CACHE = {"ts": 0.0, "data": None}
_STATS_TTL = 15.0  # 秒；避免每次刷新都遍历磁盘


def _human_size(n: float) -> str:
    """把字节数转成人类可读字符串（B/KB/MB/GB/TB）。"""
    units = ["B", "KB", "MB", "GB", "TB"]
    f = float(n)
    i = 0
    while f >= 1024 and i < len(units) - 1:
        f /= 1024.0
        i += 1
    return f"{f:.1f} {units[i]}"


def _classify(ext: str) -> str:
    """
    按扩展名归类「源文档」类型，用于概览页「类型」列与「文档类型分布」展示。

    注意：本系统 PDF/Word 导入后会额外生成一份同名的 .md（转 Markdown 产物），
    因此这里刻意区分源格式，取值范围与前端 PALETTE 的键保持一致
    （PDF / Markdown / Word / PPT / Excel / 其他），确保环形图配色命中。
    """
    e = (ext or "").lower()
    if e == ".pdf":
        return "PDF"
    if e in (".md", ".markdown", ".txt"):
        return "Markdown"
    if e in (".doc", ".docx"):
        return "Word"
    if e in (".ppt", ".pptx"):
        return "PPT"
    if e in (".xls", ".xlsx"):
        return "Excel"
    return "其他"


def compute_stats() -> dict:
    """
    计算知识库真实统计指标。
    数据来源：
      - doc/YYYYMMDD/<task_id>/ 下每个任务目录 = 一篇导入文档；
      - 任务目录内递归查找 chunks.json（数组，每项为切片）→ 已完成标记 + 切片总数；
      - 任务目录总字节数 → 存储用量；
      - 内存态 _tasks_status 中 status=="processing" 的数量 → 实时处理中。
    结果带 15s 进程内缓存，降低频繁刷新对磁盘的遍历压力。
    """
    now = time.time()
    if _STATS_CACHE["data"] is not None and now - _STATS_CACHE["ts"] < _STATS_TTL:
        return _STATS_CACHE["data"]

    root = Path(os.getenv("DATA_BASED_ROOT_DIR", r"D:\knowledge_base\doc"))
    total_docs = 0
    completed_docs = 0
    total_chunks = 0
    storage_bytes = 0
    type_counter: Dict[str, int] = {}
    recent: List[dict] = []

    if root.exists():
        for date_dir in root.iterdir():
            if not date_dir.is_dir():
                continue
            for task_dir in date_dir.iterdir():
                if not task_dir.is_dir():
                    continue
                total_docs += 1
                orig_files = [p for p in task_dir.iterdir() if p.is_file()]
                dir_size = 0
                has_chunks = False
                for p in task_dir.rglob("*"):
                    if p.is_file():
                        try:
                            sz = p.stat().st_size
                        except OSError:
                            sz = 0
                        storage_bytes += sz
                        dir_size += sz
                        if p.name == "chunks.json":
                            has_chunks = True
                            try:
                                with open(p, "r", encoding="utf-8") as fh:
                                    arr = json.load(fh)
                                if isinstance(arr, list):
                                    total_chunks += len(arr)
                            except Exception:
                                pass
                if has_chunks:
                    completed_docs += 1
                # 取一个代表性原始文件作为文档名/类型。
                # 关键：优先取「源文档」（PDF/Word/PPT/Excel），其次才是 Markdown。
                # 原因：PDF/docx 导入后会额外生成同名 .md（转 Markdown 产物），
                # 若直接取目录中首个文件，往往命中派生的 .md，
                # 导致「类型」列一律显示 Markdown，看不出“什么文档被转成了 Markdown”。
                # 只有目录里确实没有源文档时（即用户直接上传的 .md），才归为 Markdown。
                rep = None
                for f in orig_files:
                    if f.suffix.lower() in (".pdf", ".doc", ".docx",
                                            ".ppt", ".pptx", ".xls", ".xlsx"):
                        rep = f
                        break
                if rep is None:
                    for f in orig_files:
                        if f.suffix.lower() in (".md", ".markdown", ".txt"):
                            rep = f
                            break
                if rep is None and orig_files:
                    rep = orig_files[0]
                doc_name = task_dir.name
                doc_ext = ""
                if rep is not None:
                    doc_name = rep.stem
                    doc_ext = rep.suffix
                tkey = _classify(doc_ext)
                type_counter[tkey] = type_counter.get(tkey, 0) + 1
                try:
                    mtime = task_dir.stat().st_mtime
                except OSError:
                    mtime = 0.0
                recent.append({
                    "name": doc_name,
                    "type": tkey,
                    "ext": (doc_ext.replace(".", "").upper() or "—"),
                    "size": _human_size(dir_size),
                    "mtime": mtime,
                    "completed": has_chunks,
                })

    # 实时处理中任务数（来自内存态任务追踪）
    processing = 0
    try:
        from utils.task_utils import _tasks_status
        processing = sum(1 for s in _tasks_status.values() if s == "processing")
    except Exception:
        processing = 0

    recent.sort(key=lambda x: x["mtime"], reverse=True)
    recent = recent[:6]

    data = {
        "total_docs": total_docs,
        "completed_docs": completed_docs,
        "total_chunks": total_chunks,
        "storage": _human_size(storage_bytes),
        "storage_bytes": storage_bytes,
        "processing": processing,
        "type_distribution": type_counter,
        "recent": recent,
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    _STATS_CACHE["ts"] = now
    _STATS_CACHE["data"] = data
    return data


@app.get("/stats", summary="知识库统计", description="聚合 doc/ 目录与任务状态，返回文档数/切片数/存储/分类分布等真实指标")
async def stats():
    """知识库概览统计接口（供前端概览页异步拉取，替换写死的 mock 数据）。"""
    try:
        return compute_stats()
    except Exception as e:
        logger.error(f"/stats 计算失败：{str(e)}", exc_info=True)
        return {"error": str(e)}


# 3. 静态页面路由：返回文件导入前端页面
# 访问地址：http://localhost:8000/import.html
@app.get("/import.html")  # 对外访问地址
async def get_import_page():
    # 拼接HTML文件绝对路径
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "import.html"
    # 如果不存在，抛出404异常
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    # [修复] 禁止浏览器缓存页面：避免修改前端后仍加载旧版 HTML/内联 JS，
    # 出现「改了代码但页面行为没变」的错觉（本次「失败」无法定位就与此有关）。
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


# 3.1 其余静态页面路由：与导入服务同源托管，保证侧边栏跨页导航可达
@app.get("/chat.html")  # 对外访问地址
async def get_chat_page():
    # 拼接HTML文件绝对路径
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "chat.html"
    # 如果不存在，抛出404异常
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


@app.get("/overview.html")  # 对外访问地址
async def get_overview_page():
    # 拼接HTML文件绝对路径
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "overview.html"
    # 如果不存在，抛出404异常
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


# 3.1.2 登录页路由：独立托管登录页面（不挂载应用壳层，作为受保护页面的统一入口）
@app.get("/login.html")  # 对外访问地址
async def get_login_page():
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "login.html"
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


# 3.1.3 注册页路由：与登录页同源托管，复用登录页样式，形成「登录 / 注册」闭环
@app.get("/register.html")  # 对外访问地址
async def get_register_page():
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "register.html"
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


# 3.2 平台门户着陆页：访问根路径 / 或 /index.html 均返回，作为统一入口
@app.get("/")  # 门户根路径
@app.get("/index.html")  # 门户显式路径
async def get_portal_page():
    # 拼接HTML文件绝对路径
    current_dir_parent_path = Path(__file__).absolute().parent.parent
    html_path = current_dir_parent_path / "page" / "index.html"
    # 如果不存在，抛出404异常
    if not html_path.exists():
        raise HTTPException(status_code=404, detail=f"没有查询到页面，地址为：{html_path}")
    # 同样禁止缓存，保证前端改动（含主题切换脚本）即时生效
    return FileResponse(html_path, headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


# 4. 后台任务：LangGraph全流程执行
# 独立于主请求线程，由BackgroundTasks触发，避免阻塞接口响应

def convert_docx_to_md(src_path: str, dst_path: str) -> bool:
    """
    轻量 docx -> Markdown 转换（仅文本/标题/列表/表格，不走重型视觉解析）。
    用于让 Word 文档复用现有 Markdown 导入管线，快速进入切分与向量化。
    依赖 python-docx；未安装时返回 False，由上层回退为失败态。
    """
    try:
        from docx import Document
    except ImportError:
        logger.error("convert_docx_to_md 失败：未安装 python-docx，请执行 pip install python-docx")
        return False
    try:
        doc = Document(src_path)
        lines: list = []
        # 段落：按样式映射为 Markdown 标题/列表/正文
        for p in doc.paragraphs:
            style = (p.style.name or "") if getattr(p, "style", None) else ""
            text = (p.text or "").strip()
            if not text:
                lines.append("")
                continue
            if style.startswith("Title") or style.startswith("Heading 1"):
                lines.append("# " + text)
            elif style.startswith("Heading 2"):
                lines.append("## " + text)
            elif style.startswith("Heading 3"):
                lines.append("### " + text)
            elif "List" in style or style == "List Paragraph":
                lines.append("- " + text)
            else:
                lines.append(text)
        # 表格：转为 Markdown 表格
        for tbl in doc.tables:
            rows = list(tbl.rows)
            if not rows:
                continue
            n_cols = len(rows[0].cells)

            def row_to_md(r):
                return "| " + " | ".join(c.text.strip().replace("\n", " ") for c in r.cells) + " |"

            lines.append("")
            lines.append(row_to_md(rows[0]))
            lines.append("| " + " | ".join(["---"] * n_cols) + " |")
            for r in rows[1:]:
                lines.append(row_to_md(r))
            lines.append("")
        Path(dst_path).parent.mkdir(parents=True, exist_ok=True)
        with open(dst_path, "w", encoding="utf-8") as fh:
            fh.write("\n".join(lines))
        logger.info(f"docx 已转换为 Markdown：{dst_path}")
        return True
    except Exception as e:
        logger.error(f"docx 转换 Markdown 失败：{str(e)}", exc_info=True)
        return False


def run_graph_task(task_id: str, file_dir: str, import_file_path: str):
    """
    LangGraph全流程执行后台任务
    核心流程：初始化状态 → 流式执行图节点 → 实时更新任务状态 → 异常捕获
    任务状态更新：pending → processing → completed/failed
    节点进度更新：每完成一个节点，将节点名加入done_list，供前端轮询查看

    :param task_id: 全局唯一任务ID，关联单个文件的全流程处理
    :param file_dir: 该任务的本地文件存储目录（含临时文件/解析结果）
    :param import_file_path: 上传文件的本地绝对路径
    """
    try:
        # 1. 更新任务全局状态为：处理中
        update_task_status(task_id, "processing")

        # 0. docx/doc 轻量预处理：转 Markdown 后走现有 MD 管线（快，不触发重型视觉解析）
        processed_path = import_file_path
        _ext = Path(import_file_path).suffix.lower()
        if _ext in (".docx", ".doc"):
            _md_path = Path(file_dir) / (Path(import_file_path).stem + ".md")
            if convert_docx_to_md(import_file_path, str(_md_path)):
                processed_path = str(_md_path)
                add_done_task(task_id, "docx_to_md")
            else:
                update_task_status(task_id, "failed")
                set_task_result(task_id, "error", "docx 转换失败（请确认已安装 python-docx 且文件未损坏）")
                return

        # 2. 初始化LangGraph状态
        init_state = {
            "task_id": task_id, # 任务ID
            "file_dir": file_dir, # 输出路径
            "import_file_path": processed_path, # 导入路径（docx 已转为 md）
        }

        # 3. 流式执行LangGraph全流程（stream模式：实时获取每个节点的执行结果）
        workflow = KBImportWorkflow() # 创建工作流 调你整个 导入流程的 工作流
        for event in workflow.run(init_state, stream=True):# 让整个导入工作流 跑起来 流式输出
            for node_name, node_result in event.items():
                # 将完成的节点名加入【已完成列表】，前端轮询/status/{task_id}可实时获取
                add_done_task(task_id, node_name)

        # 4. 全流程执行完成，更新任务全局状态为：已完成
        update_task_status(task_id, "completed")

    except Exception as e:
        # 5. 捕获全流程异常，更新任务全局状态为：失败，并记录错误日志（含堆栈）
        update_task_status(task_id, "failed")
        # 关键：把失败原因写入任务结果，供 /status 接口回传给前端（不再只显示笼统的「失败」）
        set_task_result(task_id, "error", str(e))
        logger.info(f"[{task_id}] LangGraph全流程执行失败，异常信息：{str(e)}", exc_info=True)


# 5. 核心接口：文件上传接口
# 支持多文件上传，核心流程：接收文件 → 本地保存 → MinIO上传 → 启动后台任务
# 访问地址：http://localhost:8000/upload （POST请求，form-data格式传参）
@app.post("/upload", summary="文件上传接口", description="支持多文件批量上传，自动触发知识库导入全流程")
async def upload_files(background_tasks: BackgroundTasks, files: List[UploadFile] = File(...)):
    """
    文件上传核心接口
    1. 接收前端上传的多文件（PDF/MD为主）
    2. 按「日期/任务ID」分层保存到本地输出目录，避免文件冲突
    3. 将文件上传至MinIO对象存储，做持久化保存
    4. 为每个文件生成唯一TaskID，启动独立的LangGraph后台处理任务
    5. 实时更新任务状态，供前端轮询监控进度

    :param background_tasks: FastAPI后台任务对象，用于异步执行LangGraph流程
    :param files: 前端上传的文件列表（form-data格式）
    :return: 包含上传结果和所有任务ID的JSON响应
    """
    # 1. 构建本地存储根目录：项目根目录/doc/YYYYMMDD（按日期分层，方便管理）
    # 优先读取环境变量，未配置则使用项目内默认中转目录
    data_based_root_dir = os.getenv("DATA_BASED_ROOT_DIR", r"D:\knowledge_base\doc")
    data_dir = os.path.join(data_based_root_dir, datetime.now().strftime("%Y%m%d"))
    # 初始化任务ID列表，用于返回给前端（一个文件对应一个TaskID）
    task_ids = []

    # 2. 遍历处理每个上传的文件（多文件批量处理，各自独立生成TaskID）
    for file in files:
        # 生成全局唯一TaskID（UUID4），作为单个文件的全流程标识
        task_id = str(uuid.uuid4())# 会为每一个 PDF上传的流程 生成唯一的任务ID （追踪）
        task_ids.append(task_id)
        logger.info(f"[{task_id}] 开始处理上传文件，文件名：{file.filename}，文件类型：{file.content_type}")

        # 3. 标记「文件上传」阶段为「运行中」，前端轮询可查   添加一个标记
        add_running_task(task_id, "upload_file")# 什么都没做 就是标记一下当前的进度

        # 4. 构建该任务的本地独立目录：output/YYYYMMDD/TaskID，避免多文件重名冲突
        file_dir = os.path.join(data_dir, task_id)
        os.makedirs(file_dir, exist_ok=True)  # 目录不存在则创建，存在则不做处理
        # [安全加固 Phase0] 净化上传文件名，杜绝路径穿越（../ 或绝对路径）
        safe_name = _safe_filename(file.filename)
        # 构建上传文件的本地保存绝对路径
        import_file_path = os.path.join(file_dir, safe_name)

        # 5. 将上传的文件保存到本地临时目录（后续MinIO上传/文件解析均基于此文件）
        with open(import_file_path, "wb") as file_buffer:
            shutil.copyfileobj(file.file, file_buffer)# 复制
        logger.info(f"[{task_id}] 文件已保存至本地，路径：{import_file_path}")

        # 以上解决一个问题 把用户上传的文件 先缓存到 指定目录
        # 紧接着 进行下一步  缓存目录  用户上传的文件  上传至MinIO

        # 6. 将本地文件上传至MinIO对象存储，做持久化保存
        # 构建MinIO中的文件对象名：pdf_files/YYYYMMDD/文件名（按日期分层，和本地一致） MINIO上的路径
        # [安全加固 Phase0] 同样使用净化后的文件名，避免对象存储内路径穿越/覆盖
        minio_object_name = f"pdf_files/{datetime.now().strftime('%Y%m%d')}/{safe_name}"
        try:
            # 获取MinIO客户端实例
            minio_client = get_minio_client()

            # 从环境变量获取MinIO的桶名配置
            minio_bucket_name = minio_config.bucket_name

            # 本地文件上传至MinIO（同名文件会自动覆盖，保证文件最新）
            minio_client.fput_object(
                bucket_name=minio_bucket_name,# 桶
                object_name=minio_object_name,# 文件名 'pdf_files/20260728/H3C LA2608室内无线网关 用户手册-6W100-整本手册.pdf'
                file_path=import_file_path,# 导入路径 D:\file\20260728\d41d7662-506e-4732-8ae1-7d5613a91ffb
                content_type=file.content_type  # 传递文件原始MIME类型
            )
            logger.info(f"[{task_id}] 文件已成功上传至MinIO，桶名：{minio_bucket_name}，对象名：{minio_object_name}")
        except Exception as e:
            # MinIO上传失败，记录警告日志（不中断后续流程，本地文件仍可继续处理）
            logger.warning(f"[{task_id}] 文件上传MinIO失败，将继续执行本地处理流程，异常信息：{str(e)}", exc_info=True)

        # 先把文件接收到 本地缓存  从本地缓存里面 上传 minio
        # 7. 标记「文件上传」阶段为「已完成」，前端轮询可查
        add_done_task(task_id, "upload_file") # 标记当前进展到哪一步

        # 8. 将LangGraph全流程处理加入FastAPI后台任务（异步执行，不阻塞当前接口响应）
        background_tasks.add_task(run_graph_task, task_id, file_dir, import_file_path) # 添加后台 任务
        logger.info(f"[{task_id}] 已将LangGraph全流程加入后台任务，任务已启动")

    # 9. 所有文件处理完毕，返回上传成功信息和所有TaskID（前端基于TaskID轮询进度）
    logger.info(f"多文件上传处理完毕，共处理{len(files)}个文件，生成TaskID列表：{task_ids}")
    return {
        "code": 200,
        "message": f" 文件上传成功, total: {len(files)}",
        "task_ids": task_ids
    }


# 6. 核心接口：任务状态查询接口
# 前端轮询此接口获取单个任务的处理进度和状态
# 访问地址：http://localhost:8000/status/{task_id} （GET请求）
@app.get("/status/{task_id}", summary="任务状态查询", description="根据TaskID查询单个文件的处理进度和全局状态")
async def get_task_progress(task_id: str):
    """
    任务状态查询接口
    前端轮询此接口（如每秒1次），获取任务的实时处理进度
    返回数据均来自内存中的任务管理字典（task_utils.py），高性能无IO

    :param task_id: 全局唯一任务ID（由/upload接口返回）
    :return: 包含任务全局状态、已完成节点、运行中节点的JSON响应
    """
    # 构造任务状态返回体  获取你当前的进度
    task_status_info: Dict[str, Any] = {
        "code": 200,
        "task_id": task_id,
        "status": get_task_status(task_id),  # 任务全局状态：pending/processing/completed/failed
        "done_list": get_done_task_list(task_id),  # 已完成的节点/阶段列表
        "running_list": get_running_task_list(task_id),  # 正在运行的节点/阶段列表
        "error": get_task_result(task_id, "error", "")  # 失败时的具体原因，供前端展示定位
    }
    # 记录状态查询日志，方便追踪前端轮询情况
    logger.info(
        f"[{task_id}] 任务状态查询，当前状态：{task_status_info['status']}，已完成节点：{task_status_info['done_list']}")
    return task_status_info


# 7. 静态资源挂载：前端独立的 CSS/JS 资产以相对路径 assets/... 引用，
#    页面与资源同处 web/page 下（Vite 风格的 dist 布局），既支持 FastAPI 托管，
#    也支持 file:// 直接打开或任意 docroot 下的静态服务器，彻底避免绝对路径 404。
app.mount("/assets", StaticFiles(directory=str(Path(__file__).absolute().parent.parent / "page" / "assets")), name="assets")


if __name__ == "__main__":
    """服务启动入口：本地开发环境直接运行"""
    logger.info("鹏至智库服务启动中...")
    # 端口与 main.py 保持一致：优先读取 .env 中的 KB_PORT_IMPORT（默认 8000，
    # 当前因 8000 被「清衡舆情分析平台」占用而改为 8100），避免直接运行与一键启动端口不一致
    _import_port = int(os.getenv("KB_PORT_IMPORT", 8000))
    # 启动uvicorn服务，绑定本地IP，关闭自动重载（生产环境建议用workers多进程）
    uvicorn.run(
        app=app,
        host="127.0.0.1",  # 仅本地访问，生产环境改为0.0.0.0（允许所有IP访问）
        port=_import_port  # 服务端口（来自环境变量，默认 8000）
    )
